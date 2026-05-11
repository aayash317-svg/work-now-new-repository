import collections.abc
import os
from pptx import Presentation
from pptx.util import Inches, Pt

def create_presentation():
    prs = Presentation()

    def add_slide(prs, layout_idx, title, content=None, image_path=None):
        slide = prs.slides.add_slide(prs.slide_layouts[layout_idx])
        title_shape = slide.shapes.title
        title_shape.text = title
        if content and len(slide.placeholders) > 1:
            body_shape = slide.placeholders[1]
            tf = body_shape.text_frame
            if isinstance(content, list):
                for i, item in enumerate(content):
                    if i == 0:
                        tf.text = item
                    else:
                        p = tf.add_paragraph()
                        p.text = item
                        p.level = 0
            else:
                tf.text = content
            
            # If we are adding an image, shrink the text box width and put the image on the right
            if image_path and os.path.exists(image_path):
                body_shape.width = Inches(5.0)
                slide.shapes.add_picture(image_path, Inches(5.5), Inches(1.5), width=Inches(3.5))
                
        return slide

    # Slide 1: Title
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "WorkNow: Features & System Working"
    subtitle.text = "Deep Dive into Core Capabilities & Workflows"
    
    splash_image = r"C:\Users\ADMIN\.gemini\antigravity\brain\5522b47a-3523-4326-a85b-d9cbe2103e09\splash_screen_1777436870996.png"
    if os.path.exists(splash_image):
        slide.shapes.add_picture(splash_image, Inches(3.5), Inches(4.5), height=Inches(2.5))

    # Slide 2: Core Platform Features
    add_slide(prs, 1, "Core Platform Features", [
        "WorkNow is designed for speed, security, and seamless gig matching.",
        "Phone-based OTP Authentication (Fast login)",
        "Role-Based Dashboards (Worker vs. Employer)",
        "Real-Time Location Matching (5km radius searches)",
        "Integrated Wallet & Escrow Payments",
        "Automated Push Notifications (FCM) & SMS Alerts"
    ])

    # Slide 3: Worker Features
    map_image = r"C:\Users\ADMIN\.gemini\antigravity\brain\5522b47a-3523-4326-a85b-d9cbe2103e09\map_ui_1777436943899.png"
    add_slide(prs, 1, "Features for Workers (Students)", [
        "Map-Based Discovery: See live jobs pinned on a map based on GPS.",
        "One-Click Applications: Apply to local gigs instantly.",
        "Trust & KYC: Upload student ID/Aadhar to get a 'Verified' badge.",
        "Digital Wallet: Receive payments directly into the app.",
        "Rating System: Build reputation through successful task completions."
    ], image_path=map_image)

    # Slide 4: Employer Features
    add_slide(prs, 1, "Features for Employers", [
        "Rapid Job Posting: Create a gig with location, budget, and required skills.",
        "Automated Matching: System automatically notifies the 10 closest qualified workers.",
        "Applicant Review: View worker profiles, ratings, and KYC status before hiring.",
        "Secure Payments: Top-up wallet via Razorpay and pay workers securely.",
        "Task Management: Track jobs from 'Open' to 'In Progress' to 'Completed'."
    ])

    # Slide 5: System Working - Step 1 & 2
    add_slide(prs, 1, "How it Works: Posting & Matching", [
        "Step 1: Job Creation",
        " - Employer submits a gig via the Frontend.",
        " - The Job Service saves it and publishes a 'job.created' event via Pub/Sub.",
        "Step 2: Geo-Radius Matching",
        " - The Matching Service listens to the event.",
        " - It queries Redis for workers within a 5km radius with matching skills.",
        " - Returns a ranked list of up to 10 best workers."
    ])

    # Slide 6: System Working - Step 3 & 4
    add_slide(prs, 1, "How it Works: Applying & Hiring", [
        "Step 3: Notification & Application",
        " - Notification Service pushes FCM alerts to the matched workers.",
        " - Worker clicks the alert, views the map, and applies.",
        " - The Employer receives a push notification about the new applicant.",
        "Step 4: Hiring & Escrow",
        " - Employer selects a worker.",
        " - System conceptually escrows the funds from Employer's wallet.",
        " - Job status updates to 'IN_PROGRESS'."
    ])

    # Slide 7: System Working - Step 5
    wallet_image = r"C:\Users\ADMIN\.gemini\antigravity\brain\5522b47a-3523-4326-a85b-d9cbe2103e09\wallet_ui_1777436980123.png"
    add_slide(prs, 1, "How it Works: Completion & Payment", [
        "Step 5: Task Completion & Payout",
        " - Worker completes the task in real life.",
        " - Employer marks the job as 'Completed' in the app.",
        " - Job Service triggers 'job.completed' event.",
        " - Payment Service automatically transfers funds to the Worker's wallet.",
        " - Both parties are prompted to leave a rating."
    ], image_path=wallet_image)

    # Slide 8: Under the Hood (Backend Flow)
    add_slide(prs, 1, "Under the Hood: Technical Flow", [
        "Event-Driven Architecture ensures decoupled, scalable processing.",
        "Frontend -> API Gateway / Microservices (REST).",
        "Microservices communicate asynchronously via GCP Pub/Sub.",
        "Heavy reads (like Worker Locations) are cached in Redis.",
        "Razorpay webhooks securely update Wallet balances in PostgreSQL."
    ])

    output_path = "WorkNow_Features_And_System.pptx"
    prs.save(output_path)
    print(f"Presentation successfully saved as {output_path}")

if __name__ == "__main__":
    create_presentation()
